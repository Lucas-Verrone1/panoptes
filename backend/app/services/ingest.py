from datetime import datetime, timezone
from io import BytesIO
from pathlib import Path
import zipfile

from docx import Document as DocxDocument
from pypdf import PdfReader

from app.config import settings
from app.db import get_supabase
from app.services.chunking import chunk_text
from app.services.gemini import embed_texts, summarize_source
from app.services.n8n import notify_source_ingested
from app.services.scrape import scrape_url


def _now() -> str:
    return datetime.now(timezone.utc).isoformat()


def _trim_text(text: str, max_chars: int = settings.upload_max_text_chars) -> str:
    stripped = text.strip()
    if len(stripped) <= max_chars:
        return stripped
    return stripped[:max_chars]


def _extract_docx_text(file_bytes: bytes) -> str:
    document = DocxDocument(BytesIO(file_bytes))
    blocks: list[str] = []
    for paragraph in document.paragraphs:
        text = paragraph.text.strip()
        if text:
            blocks.append(text)
    for table in document.tables:
        for row in table.rows:
            cells = [cell.text.strip() for cell in row.cells]
            text = " | ".join(part for part in cells if part)
            if text:
                blocks.append(text)
    return _trim_text("\n\n".join(blocks))


def _extract_pdf_text(file_bytes: bytes) -> str:
    reader = PdfReader(BytesIO(file_bytes))
    pages = [page.extract_text() or "" for page in reader.pages]
    text = "\n\n".join(page.strip() for page in pages if page.strip())
    return _trim_text(text)


def _extract_text_from_uploaded_file(file_bytes: bytes, file_name: str) -> tuple[str, str, dict]:
    extension = Path(file_name).suffix.lower()

    if extension == ".pdf":
        text = _extract_pdf_text(file_bytes)
        title = Path(file_name).stem or file_name
        report = {"kind": "upload", "file_name": file_name, "format": "pdf", "characters": len(text)}
        return title, text, report

    if extension == ".docx":
        text = _extract_docx_text(file_bytes)
        title = Path(file_name).stem or file_name
        report = {"kind": "upload", "file_name": file_name, "format": "docx", "characters": len(text)}
        return title, text, report

    if extension == ".zip":
        zip_entries: list[dict[str, object]] = []
        collected_parts: list[str] = []
        with zipfile.ZipFile(BytesIO(file_bytes)) as archive:
            members = sorted(archive.infolist(), key=lambda item: item.filename)
            for member in members:
                if member.is_dir():
                    continue
                member_extension = Path(member.filename).suffix.lower()
                if member_extension not in {".pdf", ".docx", ".txt", ".md", ".html", ".htm"}:
                    continue
                try:
                    raw = archive.read(member)
                except Exception:
                    continue
                if member_extension == ".pdf":
                    extracted = _extract_pdf_text(raw)
                elif member_extension == ".docx":
                    extracted = _extract_docx_text(raw)
                else:
                    extracted = _trim_text(raw.decode("utf-8", errors="replace"))

                if not extracted:
                    continue

                if len("\n\n".join(collected_parts) + extracted) > settings.upload_max_text_chars:
                    remaining = max(0, settings.upload_max_text_chars - len("\n\n".join(collected_parts)))
                    extracted = extracted[:remaining]
                collected_parts.append(f"# {member.filename}\n\n{extracted}")
                zip_entries.append({
                    "name": member.filename,
                    "format": member_extension.lstrip("."),
                    "characters": len(extracted),
                })

        text = _trim_text("\n\n".join(collected_parts))
        if not text:
            raise ValueError("O ZIP não contém documentos legíveis com os formatos suportados.")
        title = Path(file_name).stem or file_name
        report = {
            "kind": "upload",
            "file_name": file_name,
            "format": "zip",
            "entries": zip_entries,
            "characters": len(text),
        }
        return title, text, report

    raise ValueError("Formato de arquivo não suportado.")


def _update_source(source_id: str, **fields) -> None:
    fields["updated_at"] = _now()
    get_supabase().table("sources").update(fields).eq("id", source_id).execute()


def _update_job(job_id: str, **fields) -> None:
    fields["updated_at"] = _now()
    get_supabase().table("ingestion_jobs").update(fields).eq("id", job_id).execute()


async def process_source(source_id: str, job_id: str, url: str, project_id: str) -> None:
    import asyncio
    import hashlib
    import logging
    client = get_supabase()
    claimed = client.table("ingestion_jobs").update({"status": "scraping", "updated_at": _now()}).eq("id", job_id).eq("status", "pending").execute()
    if not claimed.data:
        return
    try:
        _update_source(source_id, status="scraping", stage=2, error=None)
        scraped = await scrape_url(url)
        _update_source(source_id, extraction_report=scraped["report"])
        entries = [(chunk, page["url"]) for page in scraped["pages"] for chunk in chunk_text(page["text"])]
        chunks = [text for text, _ in entries]
        if not chunks:
            raise ValueError("Nenhum trecho útil para indexar")
        summary = await asyncio.to_thread(summarize_source, scraped["title"], scraped["text"])
        _update_source(source_id, status="embedding", stage=5)
        _update_job(job_id, status="embedding", stage=5)
        embeddings = await asyncio.to_thread(embed_texts, chunks)
        if len(embeddings) != len(chunks) or any(len(vector) != 768 for vector in embeddings):
            raise ValueError("Quantidade ou dimensão dos embeddings inválida")
        rows = [{"chunk_index": i, "content": text, "source_url": origin, "embedding": vector}
                for i, ((text, origin), vector) in enumerate(zip(entries, embeddings))]
        client.rpc("publish_source", {"p_source": source_id, "p_job": job_id,
            "p_title": scraped["title"], "p_url": scraped["url"], "p_summary": summary,
            "p_text": scraped["text"], "p_hash": hashlib.sha256(scraped["text"].encode()).hexdigest(),
            "p_report": scraped["report"], "p_chunks": rows}).execute()
    except Exception as exc:
        logging.exception("Falha na ingestão %s", job_id)
        _update_source(source_id, status="failed", error=str(exc)[:500])
        _update_job(job_id, status="failed", error=str(exc)[:500])
        return
    await deliver_pending_events()


async def process_uploaded_source(source_id: str, job_id: str, file_bytes: bytes, file_name: str, project_id: str) -> None:
    import asyncio
    import hashlib
    import logging

    client = get_supabase()
    claimed = client.table("ingestion_jobs").update({"status": "scraping", "updated_at": _now()}).eq("id", job_id).eq("status", "pending").execute()
    if not claimed.data:
        return

    try:
        _update_source(source_id, status="scraping", stage=3, error=None)
        title, text, report = _extract_text_from_uploaded_file(file_bytes, file_name)
        if not text.strip():
            raise ValueError("Nenhum conteúdo textual foi encontrado no arquivo enviado.")

        _update_source(source_id, extraction_report=report)
        entries = [(chunk, file_name) for chunk in chunk_text(text)]
        chunks = [chunk_text_chunk for chunk_text_chunk, _ in entries]
        if not chunks:
            raise ValueError("Nenhum trecho útil para indexar")

        summary = await asyncio.to_thread(summarize_source, title, text)
        _update_source(source_id, status="embedding", stage=5)
        _update_job(job_id, status="embedding", stage=5)

        embeddings = await asyncio.to_thread(embed_texts, chunks)
        if len(embeddings) != len(chunks) or any(len(vector) != 768 for vector in embeddings):
            raise ValueError("Quantidade ou dimensão dos embeddings inválida")

        rows = [{"chunk_index": i, "content": chunk, "source_url": source, "embedding": vector}
                for i, ((chunk, source), vector) in enumerate(zip(entries, embeddings))]

        client.rpc(
            "publish_source",
            {
                "p_source": source_id,
                "p_job": job_id,
                "p_title": title,
                "p_url": file_name,
                "p_summary": summary,
                "p_text": text,
                "p_hash": hashlib.sha256(text.encode()).hexdigest(),
                "p_report": report,
                "p_chunks": rows,
            },
        ).execute()
    except Exception as exc:
        logging.exception("Falha na ingestão de arquivo %s", job_id)
        _update_source(source_id, status="failed", error=str(exc)[:500])
        _update_job(job_id, status="failed", error=str(exc)[:500])
        return

    await deliver_pending_events()


async def deliver_pending_events() -> None:
    import logging
    from app.config import settings
    if not settings.n8n_webhook_url:
        return
    client = get_supabase()
    events = client.table("ingestion_events").select("*").is_("delivered_at", "null").limit(100).execute().data or []
    for event in events:
        fields = {"attempts": event["attempts"] + 1}
        try:
            await notify_source_ingested({**event["payload"], "eventId": event["id"]})
            fields.update(delivered_at=_now(), last_error=None)
        except Exception as exc:
            fields["last_error"] = str(exc)[:500]
            logging.exception("Falha ao entregar evento %s", event["id"])
        client.table("ingestion_events").update(fields).eq("id", event["id"]).execute()


async def recover_pending_sources() -> None:
    result = (
        get_supabase()
        .table("ingestion_jobs")
        .select("id,source_id,project_id,status,sources(url)")
        .eq("status", "pending")
        .execute()
    )
    for job in result.data or []:
        source = job.get("sources") or {}
        if source.get("url"):
            await process_source(job["source_id"], job["id"], source["url"], job["project_id"])

    await deliver_pending_events()
